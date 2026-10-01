"""
Parse the MCA "Struck Off Report" monthly workbooks — a different, newer format
than the one `parse_struck_off.py` handles (that script targets a single
Master_Struck_Off_Companies_and_LLPs workbook with two fixed sheets). Each of
these files covers ONE month and ONE kind (company or LLP), one sheet named
"Sheet1", with a title row identifying which ("Company Struck Off Report" /
"LLP Struck Off Report") and a header row (varying position) of
S.NO | CIN|LLPIN | Company/LLP Name.

Output columns (COPY order matches `mca_struck_off`): identifier, name, kind,
month, core_norm.

Usage:
  python parse_struck_off_monthly.py <file1.xlsx> [file2.xlsx ...] <out.csv>
"""
import openpyxl, csv, re, os, sys

SUFFIX_TOKENS = sorted([
    "limitedliabilitypartnership", "onepersoncompany", "producercompany",
    "privatelimited", "publiclimited", "companylimited", "nidhilimited",
    "privateltd", "pvtlimited", "pvtltd", "section8",
    "private", "public", "limited", "company", "nidhi",
    "llp", "opc", "llc", "ltd", "pvt",
], key=len, reverse=True)

norm = lambda s: re.sub(r"[^a-z0-9]", "", s.lower())

def core(name: str) -> str:
    s = norm(name)
    changed = True
    while changed and len(s) >= 3:
        changed = False
        for tok in SUFFIX_TOKENS:
            if len(s) - len(tok) >= 3 and s.endswith(tok):
                s = s[: -len(tok)]
                changed = True
                break
    return s

def hdr_key(v) -> str:
    return re.sub(r"[\s_]+", " ", str(v).strip().upper()) if v is not None else ""

MONTH_RE = re.compile(r"Date:\s*(\d{2}-[A-Za-z]{3}-\d{4})")

def find_header(buf):
    for r, row in enumerate(buf):
        keys = [hdr_key(c) for c in row]
        has_sno = any(k == "S.NO" or k == "SNO" for k in keys)
        has_id = any(k in ("CIN", "LLPIN") for k in keys)
        has_name = any("NAME" in k for k in keys)
        if has_id and has_name:
            id_col = next(i for i, k in enumerate(keys) if k in ("CIN", "LLPIN"))
            name_col = next(i for i, k in enumerate(keys) if "NAME" in k)
            kind = "company" if keys[id_col] == "CIN" else "llp"
            return r + 1, id_col, name_col, kind
    return None, None, None, None

def process(files, out_path):
    seen = set()
    total = 0
    with open(out_path, "w", newline="", encoding="utf-8") as fout:
        w = csv.writer(fout)
        for path in files:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            ws = wb[wb.sheetnames[0]]
            rows_iter = ws.iter_rows(values_only=True)
            buf = []
            for _ in range(20):
                try:
                    buf.append(next(rows_iter))
                except StopIteration:
                    break
            title = str(buf[0][0]) if buf and buf[0] and buf[0][0] else ""
            m = MONTH_RE.search(title)
            report_date = m.group(1) if m else ""
            hdr_idx, id_col, name_col, kind = find_header(buf)
            if hdr_idx is None:
                print(f"  SKIP (no header found): {os.path.basename(path)}")
                wb.close()
                continue
            data_rows = buf[hdr_idx:] + list(rows_iter)
            n = 0
            for row in data_rows:
                if row is None or name_col >= len(row) or not row[name_col]:
                    continue
                name = str(row[name_col]).strip().rstrip(".")
                if len(name) < 2:
                    continue
                nn = norm(name)
                if not nn or nn in seen:
                    continue
                seen.add(nn)
                ident = str(row[id_col]).strip() if id_col < len(row) and row[id_col] else ""
                w.writerow([ident, name, kind, report_date, core(name)])
                n += 1
                total += 1
            wb.close()
            print(f"  {os.path.basename(path)}: kind={kind} month={report_date} rows={n}")
    print("TOTAL:", total, "->", out_path, "size(MB):", round(os.path.getsize(out_path) / 1e6, 2))

if __name__ == "__main__":
    args = sys.argv[1:]
    if len(args) < 2:
        print(__doc__)
        sys.exit(1)
    process(args[:-1], args[-1])
